from __future__ import annotations

import shutil
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph


ROOT = Path(r"D:\UAV")
SOURCE = ROOT / "docs/research/UAV_真实遥测_可解释层次化异常检测_研究方案.docx"
OUTPUT = ROOT / "docs/research/UAV_真实遥测_层次化运行时软件故障定位_研究方案_路线C.docx"


def insert_after(paragraph: Paragraph, text: str, style: str | None = None) -> Paragraph:
    element = OxmlElement("w:p")
    paragraph._p.addnext(element)
    created = Paragraph(element, paragraph._parent)
    if style:
        created.style = style
    created.add_run(text)
    return created


def find_paragraph(document: Document, exact: str) -> Paragraph:
    for paragraph in document.paragraphs:
        if paragraph.text == exact:
            return paragraph
    raise KeyError(exact)


def replace_text(document: Document, old: str, new: str) -> None:
    paragraph = find_paragraph(document, old)
    paragraph.text = new


def add_table_row(table, values: list[str]) -> None:
    cells = table.add_row().cells
    for cell, value in zip(cells, values):
        cell.text = value


def stabilize_table_rows(table) -> None:
    for index, row in enumerate(table.rows):
        tr_pr = row._tr.get_or_add_trPr()
        cant_split = OxmlElement("w:cantSplit")
        tr_pr.append(cant_split)
        if index == 0:
            header = OxmlElement("w:tblHeader")
            header.set(qn("w:val"), "true")
            tr_pr.append(header)


shutil.copy2(SOURCE, OUTPUT)
doc = Document(OUTPUT)

# Title and research positioning.
doc.paragraphs[0].text = "面向真实无人机遥测的可解释层次化运行时软件故障定位与子系统诊断"
doc.paragraphs[1].text = "研究方案（路线 C 修订版，面向 MDPI Drones 投稿）"

replace_text(
    doc,
    "一句话概括：如何利用真实 PX4 多源飞行遥测，在准确检测无人机运行异常的基础上，进一步判断异常最可能来源于哪个状态估计/飞行子系统，并给出可解释的遥测证据？",
    "一句话概括：如何利用真实 PX4 多源飞行遥测，在检测运行异常和诊断故障域的基础上，沿 uORB 运行时接口进一步排序可疑 PX4 软件模块，并给出可追溯、可定量验证的故障定位证据？",
)
replace_text(
    doc,
    "本研究不将任务表述为严格的“源代码级软件故障定位”，而定位为运行时异常检测与子系统级诊断。原因是当前最适合复现的公开数据 UAV-SEAD 提供的是状态估计异常及异常类别标注，而不是 PX4 文件、函数或代码行级的 bug ground truth。采用这一表述可以保证研究问题与数据真实能力一致，避免将传感器、定位或状态估计异常误称为软件源代码 bug。",
    "本研究将任务定位为面向软件密集型无人机系统的遥测驱动运行时故障定位，而不是传统语句级软件缺陷定位。UAV-SEAD 用于验证真实飞行中的异常检测、故障域诊断和遥测证据；新增 PX4 固件溯源与 uORB topic→module 映射用于生成软件模块可疑度；新增 SITL 已知注入故障用于提供模块级 ground truth。只有受控注入实验可以支撑“软件模块定位”主张，真实数据中的机械、传感器和状态估计异常不得直接称为源代码 bug。",
)
replace_text(
    doc,
    "目标输出由传统的 Normal/Anomaly 二分类扩展为三级结果：第一层判断是否异常；第二层对异常窗口诊断为 Mechanical & Electrical、External Position、Global Position 或 Altitude；第三层通过 SHAP 与遥测分组给出“哪些信号、哪些 PX4 topic、哪个子系统支持该诊断”的证据链。",
    "目标输出扩展为四级结果：第一层判断是否异常；第二层诊断 Mechanical & Electrical、External Position、Global Position 或 Altitude 故障域；第三层定位高贡献 telemetry channel 与 uORB topic；第四层结合对应固件版本的软件架构图，输出可疑 PX4 module/source directory 排名。",
)

# Fill previously empty method subsections.
p = find_paragraph(doc, "4.1 数据与标签设计")
p = insert_after(p, "真实飞行分支沿用 P1–P7 的固定 UAV-SEAD 标注、10 Hz 对齐、96/8/12 窗口及原有划分，不回写既有结果。该分支的标签只支持异常窗口和故障域。软件定位分支使用 PX4 SITL 受控注入数据，额外记录 firmware commit、注入模块、源文件、注入时间与故障类型，以提供模块级定位真值。", "Normal")
p = insert_after(p, "Mechanical/Electrical 类在真实飞行评测中保留为重要故障域，但在软件模块定位评测中作为非软件来源对照；除非对应案例具有明确的软件注入记录，否则不得计为软件缺陷样本。", "Normal")

p = find_paragraph(doc, "4.2 基于哪些开源项目修改")
p = insert_after(p, "在既有 AeroTSBoost 与 ulog_annotation_tool 基础上，引入与 ULog 中 ver_sw 对应的 PX4-Autopilot 源码。静态解析 Publication、Subscription、ORB_ID、orb_advertise 与 orb_publish 等接口，建立 topic→publisher/subscriber→module→source directory 映射；SITL 采用可复现任务脚本和独立 mutation patch，不修改 P1–P7 原始数据。", "Normal")

# Add Route C methods before the experiment section.
p = find_paragraph(doc, "为避免解释仅停留在可视化，设计两个定量指标：① Top-k Feature Masking：删除 SHAP 排名前 k 的特征后重新评估，与随机删除同等数量特征比较性能下降；② Subsystem Consistency@K：统计 Top-K 高贡献信号所属 subsystem 与真实异常类别的一致程度。")
p = insert_after(p, "4.4 软件架构感知的模块可疑度", "Heading 2")
p = insert_after(p, "对每条日志提取 ver_sw、硬件和系统版本，按固件提交选择对应 PX4 源码。将 topic 级绝对 SHAP 贡献沿软件架构二部图传播到 publisher 与 subscriber 模块，分别报告 producer-only、consumer-only 和双向传播结果。模块得分需满足贡献守恒，并通过随机映射、均匀传播和 topic-only 排名进行消融。", "Normal")
p = insert_after(p, "4.5 小型 SITL 软件故障注入", "Heading 2")
p = insert_after(p, "选择 4–6 个与现有故障域相关的 PX4 模块，设计可恢复、可审计的软件 mutation，例如 uORB 发布停止/延迟、状态更新冻结、符号或单位错误、偏置注入和饱和逻辑错误。每类执行多次正常与故障任务，固定航迹、仿真环境、随机种子、注入时刻和采集通道；以注入模块为真值评估 Top-k Recall、MRR、MAP、EXAM 与检测延迟。", "Normal")

# Add experiment subsection after the current statistics paragraph.
p = find_paragraph(doc, "建议对 Macro-F1、AUPRC 等主要指标进行多随机种子或重复划分，并报告均值±标准差；如比较多个模型，可补充 Wilcoxon signed-rank test 或 bootstrap confidence interval。若数据集原始推荐划分固定，则以重复训练 seed 为主，不人为破坏 flight-level 隔离。")
p = insert_after(p, "5.3 路线 C 新增实验与验收", "Heading 2")
for text in [
    "E7 软件模块映射：统计固件版本覆盖、topic 映射率、模块数量和未解析接口；贡献传播前后必须满足数值守恒。",
    "E8 SITL 软件故障定位：至少覆盖 4 个 PX4 软件模块和 4–6 类 mutation；以 Module Top-1/3/5 Recall、MRR、MAP、EXAM 和定位延迟为主指标。",
    "E9 低成本补实验：将 direct five-class LightGBM 补齐五种子，并在 purged chronological 与 leave-log-out 下同时评估 Stage 1、Stage 2、cascade、direct five-class 及模块/子系统解释一致性。",
    "统计要求：按 flight log 或 SITL run 做 paired bootstrap 95% CI；模型随机种子标准差只描述训练波动，不替代跨航次不确定性。",
]:
    p = insert_after(p, text, "List Bullet" if text.startswith("E") else "Normal")

# Update contributions and implementation route.
replace_text(
    doc,
    "构建严格、可复现的真实飞行评测流程：主实验采用 UAV-SEAD，保留 chronological/purged/leave-log-out 防泄漏设计，并以 ALFA 做跨平台外部验证。",
    "构建真实飞行与受控软件故障互补的评测流程：UAV-SEAD 验证真实运行异常与故障域，PX4 SITL 提供软件模块真值，chronological/purged/leave-log-out 与 run-level 隔离共同控制泄漏。",
)
p = find_paragraph(doc, "构建真实飞行与受控软件故障互补的评测流程：UAV-SEAD 验证真实运行异常与故障域，PX4 SITL 提供软件模块真值，chronological/purged/leave-log-out 与 run-level 隔离共同控制泄漏。")
insert_after(p, "提出软件架构感知的证据传播方法：利用 ULog 固件提交和 uORB 发布/订阅关系，将 topic 证据转换为 PX4 软件模块可疑度，并以已知注入模块进行排名验证。", "List Number")

replace_text(
    doc,
    "为了降低研究风险，按“必须完成 → 可选增强”的顺序推进。前四个阶段构成最小可投稿版本（MVP），后续实验用于增强论文说服力。",
    "路线 C 将 P1–P7 视为已冻结研究基础，新增 P8–P10 构成面向 Drones 的软件故障定位闭环。执行顺序为：先完成固件溯源与模块映射，再搭建 SITL mutation 基准，最后补齐五种子和严格划分端到端实验。",
)
replace_text(
    doc,
    "最小可投稿版本建议固定为：UAV-SEAD + AeroTSBoost 复现 + Hierarchical LightGBM + Subsystem Diagnosis + Hierarchical SHAP Explanation + 严格防泄漏划分。ALFA 外部验证和 Unknown Fault 属于增强项；如果 CATCH/GCAD 适配成本过高，可保留 AeroTSBoost、LightGBM、RF、LSTM-AE 与单阶段 multiclass 作为最低对比集合。",
    "路线 C 的投稿版本固定为：真实 UAV-SEAD 层次诊断 + 分层 SHAP 解释 + PX4 固件感知 topic→module 映射 + 小型 SITL 软件故障注入 + direct five-class 五种子 + 严格划分端到端评估。P5–P7 作为增强与限制分析保留，不再扩展为新的主线。",
)

# Add risk boundaries.
p = find_paragraph(doc, "投稿前再次核验 2025–2026 预印本的最终发表状态、代码地址和版本，参考文献中区分正式会议/期刊与 arXiv。")
for text in [
    "真实飞行结果只能表述为 runtime failure-domain localization 或 software-subsystem suspects；只有 SITL 已知注入结果可以使用 software-module fault localization。",
    "topic→module 关系表示运行时软件依赖和证据传播路径，不自动等价于根因；正文必须区分症状模块、依赖模块与实际注入模块。",
    "所有 mutation 必须可逆、独立存档并记录基准 commit；SITL run 必须按运行实例隔离，禁止同一任务轨迹的窗口随机拆分。",
]:
    p = insert_after(p, text, "List Bullet")

# Update summary and pipeline tables.
t = doc.tables[0]
t.cell(0, 1).text = "HS-AeroTS-FL（Hierarchical Software-Subsystem-Aware AeroTS）"
t.cell(1, 1).text = "真实 UAV 运行异常检测 + 故障域诊断 + uORB/PX4 软件模块可疑度定位"
t.cell(2, 1).text = "UAV-SEAD + AeroTSBoost + PX4-Autopilot/uORB + SITL mutations + ALFA"

pipeline = ["PX4 ULog", "异常检测", "故障域诊断", "SHAP→topic", "topic→module", "SITL 真值", "模块排名+证据"]
for cell, value in zip(doc.tables[1].rows[0].cells, pipeline):
    cell.text = value

add_table_row(doc.tables[2], ["PX4 SITL mutations", "软件定位验证", "4–6 类已知模块软件故障与正常任务", "独立训练/测试", "提供 module-level ground truth"])
add_table_row(doc.tables[2], ["PX4-Autopilot source", "软件架构映射", "按 ULog ver_sw 固定版本的 uORB 发布/订阅关系", "否", "映射 topic 到模块与源码目录"])
add_table_row(doc.tables[3], ["PX4-Autopilot / SITL", "路线 C 软件定位主干", "固件源码、uORB 图、仿真任务", "静态映射、mutation patches、模块定位指标"])
add_table_row(doc.tables[4], ["E7 软件架构映射", "topic 证据能否追溯到 PX4 模块", "producer/consumer/双向传播与随机映射", "覆盖率、守恒误差、Module Consistency"])
add_table_row(doc.tables[4], ["E8 SITL 模块定位", "能否命中实际注入软件模块", "HS-AeroTS-FL、topic-only、均匀传播、随机排序", "Top-k Recall、MRR、MAP、EXAM、延迟"])
add_table_row(doc.tables[4], ["E9 补充稳健性", "层次结构是否在严格协议下成立", "cascade vs direct five-class；purged/LLO", "Macro-F1、BA、AUPRC、bootstrap CI"])
add_table_row(doc.tables[6], ["P8", "固件溯源与 uORB topic→module 映射", "版本清单、软件架构图、模块可疑度", "必须", "映射覆盖可审计且贡献守恒"])
add_table_row(doc.tables[6], ["P9", "小型 PX4 SITL 软件故障注入", "mutation 清单、运行日志、模块定位结果", "必须", "≥4 模块、≥4 类故障并报告 Top-k/MRR"])
add_table_row(doc.tables[6], ["P10", "五种子与严格划分端到端补实验", "公平对照、purged/LLO、bootstrap CI", "必须", "cascade 优势与泛化结论可复核"])

# New research resources.
p = doc.paragraphs[-1]
for text in [
    "[10] PX4-Autopilot: open-source autopilot software and uORB middleware documentation/source code.",
    "[11] Spectrum-Based Software Fault Localization surveys: definitions, suspicious program entities and Top-k/EXAM-style evaluation.",
    "[12] Research on Drone Fault Detection Based on Failure Mode Databases. Drones, 2023, 7(8), 486.",
    "[13] Event-Triggered Collaborative Fault Diagnosis for UAV–UGV Systems. Drones, 2024, 8(7), 324.",
]:
    p = insert_after(p, text, "Normal")

doc.core_properties.title = "面向真实无人机遥测的可解释层次化运行时软件故障定位与子系统诊断"
doc.core_properties.subject = "路线 C：PX4 模块映射、SITL 软件故障注入与严格端到端验证"

for section in doc.sections:
    for paragraph in section.header.paragraphs:
        if "研究方案" in paragraph.text:
            paragraph.text = "研究方案 | UAV 遥测驱动的运行时软件故障定位"

for table in doc.tables:
    stabilize_table_rows(table)

doc.save(OUTPUT)
print(OUTPUT)
